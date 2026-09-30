"""
AI Medical Report Analyzer - Model comparison pipeline
=====================================================
Cleans 5 disease datasets, then compares 6 models on each using IDENTICAL
preprocessing and evaluation (nested stratified cross-validation).

Models   : Logistic Regression, Decision Tree, Random Forest, XGBoost, SVM, MLP
Diseases : Diabetes, Heart, Kidney, Liver, Anemia
Metrics  : Accuracy, Precision, Recall, F1, ROC-AUC, Confusion Matrix,
           Training time, Prediction time
"""
import os, time, warnings, json
import glob
import numpy as np
import pandas as pd
import joblib
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import seaborn as sns

from sklearn.base import clone
from sklearn.compose import ColumnTransformer
from sklearn.impute import SimpleImputer
from sklearn.pipeline import Pipeline
from sklearn.preprocessing import StandardScaler, OneHotEncoder
from sklearn.model_selection import StratifiedKFold, GridSearchCV
from sklearn.metrics import (accuracy_score, precision_score, recall_score,
                             f1_score, roc_auc_score, confusion_matrix, roc_curve)
from sklearn.linear_model import LogisticRegression
from sklearn.tree import DecisionTreeClassifier
from sklearn.ensemble import RandomForestClassifier
from sklearn.svm import SVC
from sklearn.neural_network import MLPClassifier
from xgboost import XGBClassifier

warnings.filterwarnings("ignore")
SEED = 42
UP = os.environ.get("UPLOADS", "datasets")
OUT = os.environ.get("OUTDIR", "outputs")
os.makedirs(f"{OUT}/plots", exist_ok=True)
os.makedirs(f"{OUT}/cleaned_data", exist_ok=True)
os.makedirs(f"{OUT}/models", exist_ok=True)


# ----------------------------------------------------------------------------
# 1. DATA LOADING + CLEANING
# ----------------------------------------------------------------------------
def _finish(name, X, y, cat_cols, notes, raw_rows):
    """Drop exact duplicates (features+target), record a cleaning summary."""
    df = X.copy()
    df["target"] = y.values
    n_before = len(df)
    df = df.drop_duplicates().reset_index(drop=True)
    y = df.pop("target").astype(int)
    info = dict(
        Disease=name, Raw_rows=raw_rows, Rows_after_dedup=len(df),
        Duplicates_removed=n_before - len(df), Features=df.shape[1],
        Missing_cells=int(df.isna().sum().sum()),
        Class_0=int((y == 0).sum()), Class_1=int((y == 1).sum()),
        Minority_pct=round(100 * min((y == 0).mean(), (y == 1).mean()), 1),
        Notes=notes)
    out = df.copy(); out["target"] = y
    out.to_csv(f"{OUT}/cleaned_data/{name.lower()}_clean.csv", index=False)
    return df, y, cat_cols, info


def load_diabetes():
    d = pd.read_csv(f"{UP}/diabetes.csv")
    raw = len(d)
    y = d.pop("Outcome")

    # 1) physiologically impossible zeros -> missing
    for c in ["Glucose", "BloodPressure", "SkinThickness", "Insulin", "BMI"]:
        d.loc[d[c] == 0, c] = np.nan

    # 2) This copy was imputed with CLASS-CONDITIONAL medians (target leakage):
    #    Insulin 102.5 -> always Outcome 0, 169.5 -> always Outcome 1;
    #    SkinThickness 27/32 and BloodPressure 74.5 follow the same pattern.
    d.loc[d["Insulin"].isin([102.5, 169.5]), "Insulin"] = np.nan
    d.loc[d["SkinThickness"].isin([27, 32]), "SkinThickness"] = np.nan
    d.loc[d["BloodPressure"] == 74.5, "BloodPressure"] = np.nan

    notes = ("Pima dataset. Source file had class-conditional median imputation (target "
             "leakage: Insulin 102.5/169.5 encoded Outcome). Those fill values (Insulin 374, "
             "SkinThickness 281, BloodPressure 16 rows) and any zeros were reset to missing "
             "and are imputed inside each CV fold. No HbA1c column.")
    return _finish("Diabetes", d, y, [], notes, raw)


def load_heart():
    h = pd.read_csv(f"{UP}/heart_disease_uci.csv")
    raw = len(h)
    y = (h["num"] > 0).astype(int)          # 0 = no disease, 1-4 = disease
    h = h.drop(columns=["id", "dataset", "num", "ca", "thal", "slope"])
    for c in ["chol", "trestbps"]:
        h.loc[h[c] == 0, c] = np.nan        # 0 is physiologically impossible
    for c in ["fbs", "exang"]:
        h[c] = h[c].map({True: 1.0, False: 0.0, "True": 1.0, "False": 0.0})
    h["sex"] = h["sex"].map({"Male": 1.0, "Female": 0.0})
    notes = ("Target binarised (num>0). Dropped ca/thal/slope (66%/53%/34% missing) and "
             "id/dataset. chol=0 & trestbps=0 treated as missing. Missing values are "
             "imputed INSIDE each CV fold (no leakage). Includes clinical features "
             "(cp, restecg, exang) that a lab report does not contain.")
    return _finish("Heart", h, y, ["cp", "restecg"], notes, raw)


def load_kidney():
    """
    UCI Chronic Kidney Disease dataset.

    Expected structure:
    24 predictor variables + 1 target = 25 columns.
    Target:
        ckd     -> 1
        notckd  -> 0
    """

    
    # Find the UCI CKD ARFF file
    files = glob.glob(
        os.path.join(UP, "*chronic_kidney_disease.arff")
    )

    if not files:
        raise FileNotFoundError(
            "Could not find the UCI Chronic Kidney Disease ARFF file "
            "in the upload directory."
        )

    path = files[0]

    # ---------------------------------------------------------
    # Read ARFF file
    # ---------------------------------------------------------
    attributes = []
    data_started = False
    rows = []

    with open(path, "r", encoding="utf-8", errors="ignore") as f:

        for raw_line in f:
            line = raw_line.strip()

            if not line:
                continue

            # Attribute definitions
            if line.lower().startswith("@attribute"):
                parts = line.split(None, 2)

                if len(parts) >= 2:
                    attributes.append(parts[1])

            # Data section
            elif line.lower() == "@data":
                data_started = True

            elif data_started and not line.startswith("%"):
                # CSV-like ARFF data
                parts = [x.strip() for x in line.split(",")]
                rows.append(parts)

    # ---------------------------------------------------------
    # Verify dataset structure
    # --------------------------------------------------------- 
    
    expected_columns = 25

    print("\nUCI CKD parser check")
    print("--------------------")
    print("Attributes found:", len(attributes))
    print("Rows found:", len(rows))

    # ---------------------------------------------------------
    # Handle malformed rows
    # ---------------------------------------------------------
    # The UCI CKD file should contain:
    # 24 predictors + 1 target = 25 fields.
    #
    # Some versions of the ARFF file contain three rows with
    # one extra empty field. We remove ONLY the extra empty
    #field rather than deleting the patient record.

    valid_rows = []
    malformed_rows = []

    for i, row in enumerate(rows):

       if len(row) == expected_columns:
          valid_rows.append(row)

       elif len(row) == expected_columns + 1:

           # Remove an extra empty field if present
           empty_positions = [
              j for j, value in enumerate(row)
              if value.strip() == ""
            ]
           if len(empty_positions) == 1:
               row_fixed = [
                   value for j, value in enumerate(row)
                   if j != empty_positions[0]
                ]
               valid_rows.append(row_fixed)
               malformed_rows.append(
                   (i, len(row), "extra empty field removed")
                )
           else:
               raise ValueError(
                 f"Kidney row {i} has 26 fields but the extra "
                 f"field is not clearly empty."
            )

       else:
           raise ValueError(
              f"Kidney row {i} has {len(row)} fields. "
              f"Expected {expected_columns}."
            )

    print("Valid/repaired rows:", len(valid_rows))
    print("Rows repaired:", len(malformed_rows))

    if malformed_rows:
        print("\nRepaired kidney rows:")
        for item in malformed_rows:
            print(
              "Row index:",
               item[0],
               "-",
               item[2]
             )

    if len(valid_rows) != 400:
        raise ValueError(
           f"Expected 400 UCI CKD records, "
           f"but obtained {len(valid_rows)}."
       ) 
    # ---------------------------------------------------------
    # Build dataframe
    # ---------------------------------------------------------
    
    if len(attributes) == expected_columns:
        columns = attributes
    else:
        columns = [
          "age",
          "bp",
          "sg",
          "al",
          "su",
          "rbc",
          "pc",
          "pcc",
          "ba",
          "bgr",
          "bu",
          "sc",
          "sod",
          "pot",
          "hemo",
          "pcv",
          "wbcc",
          "rbcc",
          "htn",
          "dm",
          "cad",
          "appet",
          "pe",
          "ane",
          "class"
     ]

    # Normalize all column names
    columns = (
      pd.Index(columns)
      .astype(str)
      .str.strip()
      .str.strip("'")
      .str.strip('"')
      .str.lower()
    )

    df = pd.DataFrame(valid_rows, columns=columns)

    # ---------------------------------------------------------
    # Clean whitespace
    # ---------------------------------------------------------
    df = df.map(
    lambda x: x.strip() if isinstance(x, str) else x
)

    # Replace missing markers
    df = df.replace(
        ["?", " ?", "? "],
        np.nan
    )

    # ---------------------------------------------------------
    # Identify target column robustly
    # ---------------------------------------------------------

    # Clean column names first
    df.columns = (
      pd.Index(df.columns)
      .astype(str)
      .str.strip()
      .str.strip("'")
      .str.strip('"')
      .str.lower()
    )
    target_candidates = [
       c for c in df.columns
       if c in {"class", "classification", "target"}
    ]

    if not target_candidates:
    # In the standard UCI CKD dataset, the target is the LAST column.
    # Since we have already verified that there are exactly
    # 24 predictors + 1 target, use the last column as a safe fallback.
      target_col = df.columns[-1]

      print(
         "\nWarning: Could not identify target by name."
         f"\nUsing last column as target: '{target_col}'"
      )
    else:
        target_col = target_candidates[0]

    print("Kidney target column:", target_col)

    y_raw = (
        df[target_col]
        .astype(str)
        .str.strip()
        .str.lower()
        .str.replace(r"\s+", "", regex=True)
    )

    y = y_raw.map({
        "ckd": 1,
        "notckd": 0
    })

    if y.isna().any():
        print(
            "\nUnexpected kidney target values:",
            sorted(y_raw[y.isna()].unique())
        )

        raise ValueError(
            "Some kidney target values could not be mapped."
        )

    X = df.drop(columns=[target_col]).copy()

    # ---------------------------------------------------------
    # Convert numeric variables
    # ---------------------------------------------------------
    numeric_cols = [
        "age",
        "bp",
        "sg",
        "al",
        "su",
        "bgr",
        "bu",
        "sc",
        "sod",
        "pot",
        "hemo",
        "pcv",
        "wbcc",
        "rbcc"
    ]

    for c in numeric_cols:
        if c in X.columns:
            X[c] = pd.to_numeric(
                X[c],
                errors="coerce"
            )

    # ---------------------------------------------------------
    # Binary categorical variables
    # ---------------------------------------------------------
    binary_maps = {
        "rbc": {
            "normal": 0,
            "abnormal": 1
        },

        "pc": {
            "normal": 0,
            "abnormal": 1
        },

        "pcc": {
            "notpresent": 0,
            "present": 1
        },

        "ba": {
            "notpresent": 0,
            "present": 1
        },

        "htn": {
            "no": 0,
            "yes": 1
        },

        "dm": {
            "no": 0,
            "yes": 1
        },

        "cad": {
            "no": 0,
            "yes": 1
        },

        "appet": {
            "poor": 0,
            "good": 1
        },

        "pe": {
            "no": 0,
            "yes": 1
        },

        "ane": {
            "no": 0,
            "yes": 1
        }
    }

    for c, mapping in binary_maps.items():

        if c in X.columns:

            X[c] = (
                X[c]
                .astype(str)
                .str.strip()
                .str.lower()
                .map(mapping)
            )

    # ---------------------------------------------------------
    # Keep the original 24 predictors
    # ---------------------------------------------------------
    if X.shape[1] != 24:

        raise ValueError(
            f"UCI CKD should contain 24 predictors, "
            f"but {X.shape[1]} were found."
        )

    # ---------------------------------------------------------
    # Print verification information
    # ---------------------------------------------------------
    print("\nUCI CKD verification")
    print("--------------------")
    print("Samples:", len(X))
    print("Features:", X.shape[1])
    print("CKD:", int((y == 1).sum()))
    print("Not CKD:", int((y == 0).sum()))

    print("\nKidney columns:")
    print(list(X.columns))

    print("\nMissing values:")
    print(X.isna().sum())

    notes = (
        "UCI Chronic Kidney Disease dataset. "
        "400 records and 24 predictors. "
        "Target: ckd=1, notckd=0. "
        "Missing values represented by '?' are converted to NaN. "
        "Numeric variables are parsed as numeric and nominal variables "
        "are binary encoded. Imputation is performed inside the "
        "cross-validation pipeline."
    )

    return _finish(
        "Kidney",
        X,
        y,
        [
            "rbc",
            "pc",
            "pcc",
            "ba",
            "htn",
            "dm",
            "cad",
            "appet",
            "pe",
            "ane"
        ],
        notes,
        len(X)
    )


def load_liver():
    cols = ["Age", "Gender", "Total_Bilirubin", "Direct_Bilirubin", "Alkaline_Phosphotase",
            "ALT_SGPT", "AST_SGOT", "Total_Proteins", "Albumin", "AG_Ratio", "Selector"]
    l = pd.read_csv(f"{UP}\Indian Liver Patient Dataset (ILPD).csv", header=None, names=cols)
    raw = len(l)                            # header-less file: all rows are patients
    y = (l.pop("Selector") == 1).astype(int)  # 1 = liver patient, 2 = non-patient
    l["Gender"] = l["Gender"].map({"Male": 1.0, "Female": 0.0})
    notes = ("File has no header row - column names assigned manually (recovers the "
             "first patient that pandas otherwise consumes as header). Labels 1/2 -> 1/0. "
             "AG_Ratio missing values imputed in-fold.")
    return _finish("Liver", l, y, [], notes, raw)


def load_anemia():
    a = pd.read_excel(f"{UP}/Raw Hematological Dataset for Anemia Analysis and Classification.xlsx")
    raw = len(a)
    a.columns = a.columns.str.strip()                 # 'Gender ', 'Age ' have trailing spaces
    y = a.pop("Decision_Class")
    a["Gender"] = a["Gender"].str.strip().map({"M": 1.0, "F": 0.0})   # ' F' -> 'F'
    a.columns = ["Gender", "Age", "Hemoglobin", "RBC", "PCV_HCT", "MCV", "MCH", "MCHC"]
    notes = "CBC-only features, all extractable from a lab report. Stripped whitespace in column names/Gender values (' F'). PCV missing value imputed in-fold."
    return _finish("Anemia", a, y, [], notes, raw)


LOADERS = [load_diabetes, load_heart, load_kidney, load_liver, load_anemia]


# ----------------------------------------------------------------------------
# 2. MODELS + SMALL HYPERPARAMETER GRIDS  (same protocol for every disease)
# ----------------------------------------------------------------------------
def make_model(name, y_train):
    pos, neg = (y_train == 1).sum(), (y_train == 0).sum()
    if name == "Logistic Regression":
        return LogisticRegression(max_iter=2000, class_weight="balanced", random_state=SEED)
    if name == "Decision Tree":
        return DecisionTreeClassifier(class_weight="balanced", random_state=SEED)
    if name == "Random Forest":
        return RandomForestClassifier(class_weight="balanced", random_state=SEED, n_jobs=1)
    if name == "XGBoost":
        return XGBClassifier(eval_metric="logloss", n_jobs=1, random_state=SEED,
                             scale_pos_weight=neg / max(pos, 1), verbosity=0)
    if name == "SVM":
        return SVC(probability=True, class_weight="balanced", random_state=SEED)
    if name == "MLP Neural Network":
        return MLPClassifier(max_iter=600, early_stopping=True, n_iter_no_change=20,
                             random_state=SEED)


GRIDS = {
    "Logistic Regression": {"clf__C": [0.1, 1, 10]},
    "Decision Tree": {"clf__max_depth": [3, 5, 8, None], "clf__min_samples_leaf": [1, 5]},
    "Random Forest": {"clf__n_estimators": [200], "clf__max_depth": [None, 8],
                      "clf__min_samples_leaf": [1, 3]},
    "XGBoost": {"clf__n_estimators": [200], "clf__max_depth": [3, 5],
                "clf__learning_rate": [0.05, 0.1]},
    "SVM": {"clf__C": [0.5, 1, 5], "clf__gamma": ["scale", 0.05]},
    "MLP Neural Network": {"clf__hidden_layer_sizes": [(32,), (64, 32)],
                           "clf__alpha": [1e-4, 1e-2]},
}
MODELS = list(GRIDS)


def build_pre(X, cat_cols):
    num_cols = [c for c in X.columns if c not in cat_cols]
    transformers = [("num", Pipeline([("imp", SimpleImputer(strategy="median")),
                                      ("sc", StandardScaler())]), num_cols)]
    if cat_cols:
        transformers.append(("cat", Pipeline([
            ("imp", SimpleImputer(strategy="most_frequent")),
            ("oh", OneHotEncoder(handle_unknown="ignore"))]), cat_cols))
    return ColumnTransformer(transformers)


def make_search(name, pre, y_train):
    est = Pipeline([("pre", clone(pre)), ("clf", make_model(name, y_train))])
    return GridSearchCV(est, GRIDS[name], cv=StratifiedKFold(3, shuffle=True, random_state=SEED),
                        scoring="roc_auc", n_jobs=1, refit=True)


# ----------------------------------------------------------------------------
# 3. NESTED CROSS-VALIDATION
# ----------------------------------------------------------------------------
def evaluate(datasets):
    fold_rows, oof, cms = [], {}, {}
    outer = StratifiedKFold(5, shuffle=True, random_state=SEED)
    for dname, (X, y, cat_cols, _) in datasets.items():
        pre = build_pre(X, cat_cols)
        for m in MODELS:
            t_start = time.time()
            oof_y, oof_p = [], []
            cm = np.zeros((2, 2), int)
            for k, (tr, te) in enumerate(outer.split(X, y), 1):
                gs = make_search(m, pre, y.iloc[tr])
                gs.fit(X.iloc[tr], y.iloc[tr])
                t0 = time.perf_counter()
                proba = gs.predict_proba(X.iloc[te])[:, 1]
                pred_ms = (time.perf_counter() - t0) / len(te) * 1000
                pred = (proba >= 0.5).astype(int)
                yt = y.iloc[te]
                cm += confusion_matrix(yt, pred, labels=[0, 1])
                oof_y.append(yt.values); oof_p.append(proba)
                fold_rows.append(dict(
                    Disease=dname, Model=m, Fold=k,
                    Accuracy=accuracy_score(yt, pred),
                    Precision=precision_score(yt, pred, zero_division=0),
                    Recall=recall_score(yt, pred, zero_division=0),
                    F1=f1_score(yt, pred, zero_division=0),
                    ROC_AUC=roc_auc_score(yt, proba),
                    Train_time_s=gs.refit_time_,          # fit time of the tuned model
                    Predict_ms_per_sample=pred_ms,
                    Best_params=json.dumps({a.replace("clf__", ""): (list(b) if isinstance(b, tuple) else b)
                                            for a, b in gs.best_params_.items()}, default=str)))
            oof[(dname, m)] = (np.concatenate(oof_y), np.concatenate(oof_p))
            cms[(dname, m)] = cm
            print(f"[{dname:9s}] {m:20s} done in {time.time()-t_start:5.1f}s", flush=True)
    return pd.DataFrame(fold_rows), oof, cms


# ----------------------------------------------------------------------------
# 4. TABLES
# ----------------------------------------------------------------------------
METRICS = ["Accuracy", "Precision", "Recall", "F1", "ROC_AUC"]


def summarise(folds):
    agg = folds.groupby(["Disease", "Model"]).agg(
        **{f"{m}_mean": (m, "mean") for m in METRICS},
        **{f"{m}_std": (m, "std") for m in METRICS},
        Train_time_s=("Train_time_s", "mean"),
        Predict_ms_per_sample=("Predict_ms_per_sample", "mean")).reset_index()
    return agg


def pivot_pm(summary, metric, decimals=3):
    p = summary.copy()
    p["v"] = p.apply(lambda r: f"{r[metric+'_mean']:.{decimals}f} ± {r[metric+'_std']:.{decimals}f}", axis=1)
    return p.pivot(index="Model", columns="Disease", values="v").reindex(MODELS)[list(DISEASE_ORDER)]


def best_models(summary):
    rows = []
    for d, g in summary.groupby("Disease"):
        g = g.sort_values(["ROC_AUC_mean", "F1_mean"], ascending=False)
        b = g.iloc[0]

        # Models within one standard deviation of the best mean ROC-AUC.
        # This is descriptive only and is NOT a statistical significance test.
        close = g[
            g.ROC_AUC_mean >= b.ROC_AUC_mean - b.ROC_AUC_std
        ]["Model"].tolist()

        rows.append(dict(
            Disease=d,
            Best_model=b.Model,
            ROC_AUC=round(b.ROC_AUC_mean, 4),
            F1=round(b.F1_mean, 4),
            Recall=round(b.Recall_mean, 4),
            Accuracy=round(b.Accuracy_mean, 4),
            Models_within_1_SD_of_best_AUC=", ".join(close)
        ))

    return pd.DataFrame(rows)


# ----------------------------------------------------------------------------
# 5. PLOTS
# ----------------------------------------------------------------------------
PAL = dict(zip(MODELS, sns.color_palette("tab10", 6)))
DISEASE_ORDER = ("Diabetes", "Heart", "Kidney", "Liver", "Anemia")


def plot_metric_bars(summary):
    fig, axes = plt.subplots(2, 3, figsize=(20, 10))
    for ax, d in zip(axes.flat, DISEASE_ORDER):
        g = summary[summary.Disease == d].set_index("Model").reindex(MODELS)
        w = 0.13
        for i, m in enumerate(MODELS):
            ax.bar(np.arange(len(METRICS)) + i * w,
                   [g.loc[m, f"{x}_mean"] for x in METRICS], w,
                   yerr=[g.loc[m, f"{x}_std"] for x in METRICS], capsize=2,
                   color=PAL[m], label=m)
        ax.set_xticks(np.arange(len(METRICS)) + 2.5 * w)
        ax.set_xticklabels([x.replace("_", "-") for x in METRICS])
        ax.set_ylim(0.4, 1.02); ax.set_title(d, fontweight="bold"); ax.grid(axis="y", alpha=.3)
    h, l = axes.flat[0].get_legend_handles_labels()
    axes.flat[5].axis("off"); axes.flat[5].legend(h, l, loc="center", fontsize=13, title="Model")
    fig.suptitle("Model comparison per disease (5-fold nested CV, mean ± SD)", fontsize=16)
    fig.tight_layout(); fig.savefig(f"{OUT}/plots/1_metric_comparison.png", dpi=150); plt.close(fig)


def plot_heatmaps(summary):
    fig, axes = plt.subplots(1, 3, figsize=(20, 5.5))
    for ax, (met, title) in zip(axes, [("ROC_AUC", "ROC-AUC"), ("F1", "F1-score"), ("Recall", "Recall")]):
        p = summary.pivot(index="Model", columns="Disease", values=f"{met}_mean").reindex(MODELS)[list(DISEASE_ORDER)]
        sns.heatmap(p, annot=True, fmt=".3f", cmap="YlGnBu", vmin=0.5, vmax=1, ax=ax, cbar=False)
        ax.set_title(title, fontweight="bold"); ax.set_ylabel("")
    fig.tight_layout(); fig.savefig(f"{OUT}/plots/2_heatmaps_auc_f1_recall.png", dpi=150); plt.close(fig)


def plot_roc(oof):
    fig, axes = plt.subplots(2, 3, figsize=(19, 11))
    for ax, d in zip(axes.flat, DISEASE_ORDER):
        for m in MODELS:
            y, p = oof[(d, m)]
            fpr, tpr, _ = roc_curve(y, p)
            ax.plot(fpr, tpr, color=PAL[m], lw=2, label=f"{m} ({roc_auc_score(y, p):.3f})")
        ax.plot([0, 1], [0, 1], "k--", lw=1)
        ax.set_title(d, fontweight="bold"); ax.set_xlabel("False positive rate"); ax.set_ylabel("True positive rate")
        ax.legend(loc="lower right", fontsize=9); ax.grid(alpha=.3)
    axes.flat[5].axis("off")
    fig.suptitle("ROC curves (pooled out-of-fold predictions)", fontsize=16)
    fig.tight_layout(); fig.savefig(f"{OUT}/plots/3_roc_curves.png", dpi=150); plt.close(fig)


def plot_confusion(cms):
    fig, axes = plt.subplots(5, 6, figsize=(21, 17))
    for i, d in enumerate(DISEASE_ORDER):
        for j, m in enumerate(MODELS):
            cm = cms[(d, m)]
            sns.heatmap(cm, annot=True, fmt="d", cmap="Blues", cbar=False, ax=axes[i, j],
                        xticklabels=["Neg", "Pos"], yticklabels=["Neg", "Pos"])
            if i == 0: axes[i, j].set_title(m, fontweight="bold")
            if j == 0: axes[i, j].set_ylabel(f"{d}\nTrue", fontweight="bold")
            else: axes[i, j].set_ylabel("")
            axes[i, j].set_xlabel("Predicted")
    fig.suptitle("Confusion matrices (summed over 5 outer folds)", fontsize=16)
    fig.tight_layout(); fig.savefig(f"{OUT}/plots/4_confusion_matrices.png", dpi=130); plt.close(fig)


def plot_times(summary):
    fig, axes = plt.subplots(1, 2, figsize=(18, 6))
    for ax, col, title in [(axes[0], "Train_time_s", "Training time (s, tuned model refit)"),
                           (axes[1], "Predict_ms_per_sample", "Prediction time (ms per sample)")]:
        p = summary.pivot(index="Disease", columns="Model", values=col).reindex(DISEASE_ORDER)[MODELS]
        p.plot(kind="bar", ax=ax, color=[PAL[m] for m in MODELS], logy=True, width=.85)
        ax.set_title(title, fontweight="bold"); ax.set_xlabel(""); ax.tick_params(axis="x", rotation=0)
        ax.grid(axis="y", alpha=.3)
    fig.tight_layout(); fig.savefig(f"{OUT}/plots/5_training_prediction_time.png", dpi=150); plt.close(fig)


# ----------------------------------------------------------------------------
# 6. EXCEL REPORT
# ----------------------------------------------------------------------------
def write_excel(cleaning, summary, folds, best):
    from openpyxl import load_workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    path = f"{OUT}/model_comparison_results.xlsx"
    with pd.ExcelWriter(path, engine="openpyxl") as w:
        cleaning.to_excel(w, sheet_name="Cleaning_Summary", index=False)
        best.to_excel(w, sheet_name="Best_Model_Per_Disease", index=False)
        for m in METRICS:
            pivot_pm(summary, m).reset_index().to_excel(w, sheet_name=m, index=False)
        t = summary[["Disease", "Model", "Train_time_s", "Predict_ms_per_sample"]].round(5)
        t.to_excel(w, sheet_name="Timing", index=False)
        summary.round(4).to_excel(w, sheet_name="Full_Summary", index=False)
        folds.round(5).to_excel(w, sheet_name="Fold_Level_Results", index=False)
        pd.DataFrame({"Notes": [
    "Values are ML outputs computed by disease_model_comparison.py (static values, no formulas).",
    "Protocol: 5-fold stratified outer CV; inner 3-fold GridSearchCV (scoring=ROC-AUC) per outer fold.",
    "Preprocessing (median/mode imputation, scaling, one-hot) is fit inside each training fold - no leakage.",
    "Class imbalance handled with class_weight='balanced' (LR/DT/RF/SVM) and scale_pos_weight (XGBoost). MLP has no weighting.",
    "Predictions use a 0.5 probability threshold. Precision/Recall/F1 are for the positive (disease) class.",
    "'±' values are standard deviation across the 5 outer folds.",
    "Models within 1 SD of the best ROC-AUC are reported descriptively; this is not a statistical significance or equivalence test.",
    "Training time = refit time of the tuned model on the training fold (excludes grid search).",
    "MLP is scikit-learn MLPClassifier (TensorFlow/Keras was not available in this environment).",
    "Kidney uses the real UCI CKD dataset (400 rows)."
]}
).to_excel(w, sheet_name="Notes", index=False)
    wb = load_workbook(path)
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                c.font = Font(name="Arial", size=10, bold=(c.row == 1),
                              color="FFFFFF" if c.row == 1 else "000000")
                if c.row == 1: c.fill = PatternFill("solid", fgColor="1F4E78")
                c.alignment = Alignment(wrap_text=(ws.title in ("Notes", "Cleaning_Summary")), vertical="top")
        for col in ws.columns:
            width = max(len(str(c.value)) if c.value is not None else 0 for c in col[:40])
            ws.column_dimensions[col[0].column_letter].width = min(max(12, width + 2), 90 if ws.title in ("Notes", "Cleaning_Summary") else 32)
        ws.freeze_panes = "A2"
    wb.save(path)


# ----------------------------------------------------------------------------
# 7. MAIN
# ----------------------------------------------------------------------------
def main():
    datasets, cleaning = {}, []
    for fn in LOADERS:
        X, y, cat, info = fn()
        name = info["Disease"]
        datasets[name] = (X, y, cat, info)
        cleaning.append(info)
        print(f"Cleaned {name}: {info['Rows_after_dedup']} rows, {info['Features']} features, "
              f"{info['Duplicates_removed']} duplicates removed", flush=True)
    cleaning = pd.DataFrame(cleaning)

    folds, oof, cms = evaluate(datasets)
    summary = summarise(folds)
    best = best_models(summary)

    folds.to_csv(f"{OUT}/fold_level_results.csv", index=False)
    summary.to_csv(f"{OUT}/summary_results.csv", index=False)
    write_excel(cleaning, summary, folds, best)

    plot_metric_bars(summary); plot_heatmaps(summary); plot_roc(oof)
    plot_confusion(cms); plot_times(summary)

    # Refit the best model per disease on ALL data and save it for the app
    for _, r in best.iterrows():
        if r.ROC_AUC < 0.6:
            print(f"Skipping model export for {r.Disease}: ROC-AUC {r.ROC_AUC} ~ chance")
            continue
        X, y, cat, _ = datasets[r.Disease]
        gs = make_search(r.Best_model, build_pre(X, cat), y)
        gs.fit(X, y)
        joblib.dump({"model": gs.best_estimator_, "features": list(X.columns),
                     "disease": r.Disease, "algorithm": r.Best_model},
                    f"{OUT}/models/{r.Disease.lower()}_{r.Best_model.lower().replace(' ', '_')}.joblib")

    print("\n=== BEST MODEL PER DISEASE ===")
    print(best.to_string(index=False))
    print("\n=== ROC-AUC (mean ± SD) ===")
    print(pivot_pm(summary, "ROC_AUC").to_string())


if __name__ == "__main__":
    main()