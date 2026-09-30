"""
Feature-availability experiment
===============================
Your app extracts values from an uploaded lab report. Some dataset features
(chest-pain type, pedal edema, skin-thickness, ...) are NOT on a lab report.
This script measures how much performance is lost when a model may only use
features that a report (or a few manually-typed vitals) can supply.

Tiers per disease
  Report-only     : values found on a lab report (labs + age/sex from header)
  Report + manual : report-only + simple vitals/history the user can type in
  Full            : every cleaned feature (baseline, read from summary_results.csv)

Liver and Anemia are already 100% lab-report features, so they are skipped.
Uses the SAME protocol as disease_model_comparison.py (nested stratified CV).
"""
import os
import pandas as pd
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import seaborn as sns
import disease_model_comparison as base   # re-use loaders / models / evaluate()

OUT = base.OUT
MODELS = base.MODELS

VARIANTS = {
    "Diabetes": {
        "Report-only": ["Glucose", "Insulin", "Age"],
        "Report + manual": ["Glucose", "Insulin", "Age", "BMI", "BloodPressure", "Pregnancies"],
    },
    "Heart": {
        "Report-only": ["age", "sex", "chol", "fbs"],
        "Report + manual": ["age", "sex","cp", "chol", "fbs", "trestbps", "thalch"],
    },
    "Kidney": {
        # urine routine + blood chemistry + CBC are all on a standard lab report
        "Report-only": ["age", "sg", "al", "su", "rbc", "pc", "pcc", "ba", "bgr", "bu", "sc",
                        "sod", "pot", "hemo", "pcv", "wbcc", "rbcc"],
        "Report + manual": ["age", "sg", "al", "su", "rbc", "pc", "pcc", "ba", "bgr", "bu", "sc",
                            "sod", "pot", "hemo", "pcv", "wbcc", "rbcc",
                            "bp", "htn", "dm", "cad"],
    },
}
LOADERS = {"Diabetes": base.load_diabetes, "Heart": base.load_heart, "Kidney": base.load_kidney}


def main():
    full = pd.read_csv(f"{OUT}/summary_results.csv")
    full = full[full.Disease.isin(VARIANTS)].assign(Variant="Full")
    n_full = {}
    datasets = {}
    for d, fn in LOADERS.items():
        X, y, cat, _ = fn()
        n_full[d] = X.shape[1]
        for vname, cols in VARIANTS[d].items():
            cols = [c for c in cols if c in X.columns]
            datasets[f"{d}||{vname}"] = (X[cols], y, [c for c in cat if c in cols], None)

    folds, oof, cms = base.evaluate(datasets)
    folds[["Disease", "Variant"]] = folds["Disease"].str.split(r"\|\|", expand=True)
    # summarise per (Disease, Variant)
    parts = []
    for v, g in folds.groupby("Variant"):
        s = base.summarise(g); s["Variant"] = v; parts.append(s)
    summ = pd.concat(parts + [full], ignore_index=True)
    summ.to_csv(f"{OUT}/feature_variant_results.csv", index=False)

    nfeat = {(d, v): len([c for c in cols]) for d, vs in VARIANTS.items() for v, cols in vs.items()}
    order = ["Report-only", "Report + manual", "Full"]

    # ---- table: ROC-AUC mean ± sd  (Disease/Variant x Model) ----
    rows = []
    for (d, v), g in summ.groupby(["Disease", "Variant"]):
        g = g.set_index("Model").reindex(MODELS)
        r = {"Disease": d, "Variant": v,
             "n_features": n_full[d] if v == "Full" else nfeat[(d, v)]}
        for m in MODELS:
            r[m] = f"{g.loc[m,'ROC_AUC_mean']:.3f} ± {g.loc[m,'ROC_AUC_std']:.3f}"
        b = g.ROC_AUC_mean.idxmax()
        r["Best"] = f"{b} ({g.loc[b,'ROC_AUC_mean']:.3f})"
        r["Best_F1"] = round(g.loc[b, "F1_mean"], 3); r["Best_Recall"] = round(g.loc[b, "Recall_mean"], 3)
        rows.append(r)
    tab = pd.DataFrame(rows)
    tab["Variant"] = pd.Categorical(tab.Variant, order, ordered=True)
    tab = tab.sort_values(["Disease", "Variant"]).reset_index(drop=True)
    tab.to_csv(f"{OUT}/feature_variant_table.csv", index=False)

    # ---- Excel ----
    from openpyxl import load_workbook
    from openpyxl.styles import Font, PatternFill
    xp = f"{OUT}/feature_variant_comparison.xlsx"
    with pd.ExcelWriter(xp, engine="openpyxl") as w:
        tab.to_excel(w, sheet_name="ROC_AUC_by_variant", index=False)
        summ.round(4).to_excel(w, sheet_name="Full_Summary", index=False)
        pd.DataFrame({"Notes": [
            "Report-only = features present on a standard lab report (plus age/sex).",
            "Report + manual = report-only + vitals/history a user can type in (BP, BMI, HR, hypertension...).",
            "Full = all cleaned dataset features (clinical inputs such as chest-pain type are not on a lab report).",
            "Same nested 5x3 CV protocol as the main comparison; values are mean ± SD over outer folds.",
            "Liver and Anemia already use only lab-report features."]}).to_excel(w, sheet_name="Notes", index=False)
    wb = load_workbook(xp)
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                c.font = Font(name="Arial", size=10, bold=(c.row == 1), color="FFFFFF" if c.row == 1 else "000000")
                if c.row == 1: c.fill = PatternFill("solid", fgColor="1F4E78")
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = min(max(12, max(len(str(c.value or "")) for c in col[:30]) + 2), 80)
        ws.freeze_panes = "A2"
    wb.save(xp)

    # ---- plot: AUC per model, grouped by variant, one panel per disease ----
    fig, axes = plt.subplots(1, 3, figsize=(20, 5.5))
    pal = dict(zip(order, ["#d95f02", "#7570b3", "#1b9e77"]))
    for ax, d in zip(axes, VARIANTS):
        for i, v in enumerate(order):
            g = summ[(summ.Disease == d) & (summ.Variant == v)].set_index("Model").reindex(MODELS)
            ax.bar([x + i * 0.27 for x in range(len(MODELS))], g.ROC_AUC_mean, 0.27,
                   yerr=g.ROC_AUC_std, capsize=2, color=pal[v], label=v)
        ax.set_xticks([x + 0.27 for x in range(len(MODELS))])
        ax.set_xticklabels([m.replace(" ", "\n") for m in MODELS], fontsize=9)
        ax.set_ylim(0.5, 1.02); ax.set_title(d, fontweight="bold"); ax.set_ylabel("ROC-AUC"); ax.grid(axis="y", alpha=.3)
    axes[0].legend(title="Feature set")
    fig.suptitle("What is lost when models may only use lab-report features? (ROC-AUC, mean ± SD)", fontsize=15)
    fig.tight_layout(); fig.savefig(f"{OUT}/plots/6_feature_variants.png", dpi=150); plt.close(fig)

    pd.set_option("display.width", 250); pd.set_option("display.max_columns", 20)
    print("\n=== ROC-AUC BY FEATURE SET ===")
    print(tab.to_string(index=False))


if __name__ == "__main__":
    main()