"""
Ablation study: Kidney and Anemia
==================================
Both models score near-ceiling AUC. This is NOT leakage (no test-fold information
touches training - preprocessing is fit inside each CV fold, as in every other
script here). It is criterion-dependence: the diagnostic label itself was assigned
using markers very close to the model's own features (anemia is clinically DEFINED
by a hemoglobin cutoff; this CKD dataset's diagnoses were made by a nephrologist
using largely the same lab panel). This script measures how much AUC depends on
those defining markers, for honest reporting in the paper - it does not "fix"
anything, because there is nothing broken to fix.
"""
import pandas as pd
import numpy as np
from sklearn.model_selection import StratifiedKFold, cross_val_predict
from sklearn.linear_model import LogisticRegression
from sklearn.metrics import roc_auc_score
from sklearn.pipeline import Pipeline
from sklearn.base import clone

import disease_model_comparison as base

OUT = base.OUT
SEED = base.SEED
cv = StratifiedKFold(5, shuffle=True, random_state=SEED)
clf = LogisticRegression(max_iter=2000, class_weight="balanced", random_state=SEED)


def auc_for(X, y, cat_cols=()):
    pre = base.build_pre(X, list(cat_cols))
    pipe = Pipeline([("pre", clone(pre)), ("clf", clone(clf))])
    proba = cross_val_predict(pipe, X, y, cv=cv, method="predict_proba", n_jobs=1)[:, 1]
    return roc_auc_score(y, proba)


def main():
    rows = []

    # ---------------- KIDNEY ----------------
    X, y, _, _ = base.load_kidney()
    report_cols = ["age", "sg", "al", "su", "rbc", "pc", "pcc", "ba", "bgr", "bu", "sc",
                  "sod", "pot", "hemo", "pcv", "wbcc", "rbcc"]
    steps = [
        ("A. Report-only (17 features)", report_cols),
        ("B. Without sg, al (urinalysis markers)", [c for c in report_cols if c not in ("sg", "al")]),
        ("C. Without sg, al, sc, hemo (core diagnostic panel)",
         [c for c in report_cols if c not in ("sg", "al", "sc", "hemo")]),
        ("D. Basic chemistry only (bgr, bu, sc, sod, pot, hemo)",
         ["bgr", "bu", "sc", "sod", "pot", "hemo"]),
    ]
    for label, cols in steps:
        a = auc_for(X[cols], y)
        rows.append(dict(Disease="Kidney", Step=label, n_features=len(cols), ROC_AUC=round(a, 4)))

    # ---------------- ANEMIA ----------------
    X, y, _, _ = base.load_anemia()
    all_cols = list(X.columns)
    steps = [
        ("A. All 8 features", all_cols),
        ("B. Without Hemoglobin", [c for c in all_cols if c != "Hemoglobin"]),
        ("C. Without Hemoglobin + PCV_HCT", [c for c in all_cols if c not in ("Hemoglobin", "PCV_HCT")]),
        ("D. RBC indices only (MCV, MCH, MCHC, RBC)", ["MCV", "MCH", "MCHC", "RBC"]),
    ]
    for label, cols in steps:
        a = auc_for(X[cols], y)
        rows.append(dict(Disease="Anemia", Step=label, n_features=len(cols), ROC_AUC=round(a, 4)))

    tab = pd.DataFrame(rows)
    tab.to_csv(f"{OUT}/ablation_results.csv", index=False)

    from openpyxl import load_workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    xp = f"{OUT}/ablation_results.xlsx"
    with pd.ExcelWriter(xp, engine="openpyxl") as w:
        tab.to_excel(w, sheet_name="Ablation", index=False)
        pd.DataFrame({"Notes": [
            "This is NOT a leakage test - preprocessing is fit inside each CV fold in every "
            "script in this project, so no test-fold information reaches training anywhere.",
            "This measures CRITERION-DEPENDENCE: how much AUC relies on the specific markers "
            "that were likely used to assign the diagnostic label in the first place.",
            "Anemia is clinically DEFINED by a hemoglobin cutoff (WHO criteria), so Hemoglobin "
            "carrying most of the signal is expected, not a modelling artefact.",
            "The UCI CKD dataset's labels were assigned by a nephrologist using largely the same "
            "lab panel (sg, al, sc, hemo are core CKD diagnostic markers), so a large AUC drop "
            "when they are removed is expected, not a bug.",
            "A large drop after removing the defining marker does not mean the model is broken; "
            "it means the model has (unsurprisingly) learned the same clinical rule used to "
            "produce the label. Report this transparently rather than treating it as an error."]}
        ).to_excel(w, sheet_name="Notes", index=False)
    wb = load_workbook(xp)
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                c.font = Font(name="Arial", size=10, bold=(c.row == 1), color="FFFFFF" if c.row == 1 else "000000")
                if c.row == 1: c.fill = PatternFill("solid", fgColor="1F4E78")
                c.alignment = Alignment(wrap_text=(ws.title == "Notes"), vertical="top")
        for col in ws.columns:
            w_ = max(len(str(c.value)) if c.value is not None else 0 for c in col[:10])
            ws.column_dimensions[col[0].column_letter].width = min(max(14, w_ + 2), 90 if ws.title == "Notes" else 45)
    wb.save(xp)

    print(tab.to_string(index=False))


if __name__ == "__main__":
    main()